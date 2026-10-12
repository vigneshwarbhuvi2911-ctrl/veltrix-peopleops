"""Format-specific parsers: PDF, DOCX, HTML and Markdown -> a flat list of Blocks.

A Block is a heading, a paragraph of text or a table, in reading order, with the page
number it starts on (PDF only). Tables keep their rows so that the chunker can render
them as Markdown and the holiday extractor can read cells directly.

Nothing here is specific to the Veltrix corpus: headings are found from document
structure (PDF font size/weight, DOCX styles, HTML tags, Markdown #), and the
"metadata card" that every document starts with is recognised by its field labels.
"""
from __future__ import annotations

import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from app.ingestion.ocr import ocr_pdf_page

log = logging.getLogger(__name__)

META_LABELS = {"document id", "version", "effective date", "applies to", "owner", "classification", "field"}
_META_LINE = re.compile(r"^(document id|version|effective date|applies to|owner|classification)\b", re.I)


@dataclass
class Block:
    kind: str  # "heading" | "text" | "table"
    text: str = ""
    level: int = 1  # headings: 0 = document title, 1+ = section levels
    page: int | None = None
    rows: list[list[str]] = field(default_factory=list)  # tables only


@dataclass
class ParsedDocument:
    blocks: list[Block]
    page_count: int | None = None
    ocr_pages: list[int] = field(default_factory=list)


def _clean(s: str | None) -> str:
    s = (s or "").replace(" ", " ")
    return re.sub(r"[ \t]+", " ", s).strip()


def _clean_cell(s: str | None) -> str:
    return re.sub(r"\s+", " ", (s or "").replace(" ", " ")).strip()


def is_metadata_table(rows: list[list[str]]) -> bool:
    """The document-control card (Document ID / Version / Effective date ...)."""
    labels: list[str] = []
    for row in rows:
        if row:
            labels.extend(l.strip().lower() for l in (row[0] or "").split("\n"))
    hits = sum(1 for l in labels if l in META_LABELS)
    return hits >= 3


def table_to_markdown(rows: list[list[str]]) -> str:
    rows = [[_clean_cell(c) for c in r] for r in rows if any(_clean_cell(c) for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


# --------------------------------------------------------------------------- PDF
_NUMBERED_HEADING = re.compile(r"^(\d+(\.\d+)*\.?|[A-Z]\.)\s+\S")
_FOOTER_PAGE = re.compile(r"\bPage\s+\d+\s*$", re.I)


def parse_pdf(path: Path, ocr_lang: str = "eng") -> ParsedDocument:
    import pdfplumber

    blocks: list[Block] = []
    ocr_pages: list[int] = []
    with pdfplumber.open(str(path)) as pdf:
        # Body font size = the most common character size in the document.
        sizes: Counter = Counter()
        for page in pdf.pages:
            for ch in page.chars:
                sizes[round(ch["size"], 1)] += 1
        body_size = sizes.most_common(1)[0][0] if sizes else 10.0
        max_size = max(sizes) if sizes else body_size

        seen_section = False
        for pno, page in enumerate(pdf.pages, start=1):
            text_chars = [c for c in page.chars if c["text"].strip()]
            if len(text_chars) < 20:
                # No usable text layer: scanned page.
                text = ocr_pdf_page(page, lang=ocr_lang)
                if text.strip():
                    ocr_pages.append(pno)
                    for para in re.split(r"\n\s*\n", text):
                        para = "\n".join(_clean(l) for l in para.splitlines() if l.strip())
                        if para:
                            blocks.append(Block("text", para, page=pno))
                continue

            tables = page.find_tables()
            items: list[tuple[float, Block]] = []
            bboxes = []
            for t in tables:
                rows = [[c or "" for c in r] for r in t.extract()]
                bboxes.append(t.bbox)
                if not rows or is_metadata_table(rows):
                    continue
                items.append((t.bbox[1], Block("table", page=pno, rows=[[_clean_cell(c) for c in r] for r in rows])))

            def in_table(w) -> bool:
                cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
                return any(x0 - 1 <= cx <= x1 + 1 and top - 1 <= cy <= bottom + 1 for x0, top, x1, bottom in bboxes)

            words = page.extract_words(extra_attrs=["fontname", "size"], use_text_flow=True)
            lines: list[list[dict]] = []
            for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
                if in_table(w):
                    continue
                if lines and abs(lines[-1][0]["top"] - w["top"]) <= 3:
                    lines[-1].append(w)
                else:
                    lines.append([w])

            para: list[str] = []
            para_top = None
            prev_bottom = None

            def flush():
                nonlocal para, para_top
                if para:
                    items.append((para_top, Block("text", _clean(" ".join(para)), page=pno)))
                para, para_top = [], None

            for ln in lines:
                ln.sort(key=lambda w: w["x0"])
                text = _clean(" ".join(w["text"] for w in ln))
                top = ln[0]["top"]
                size = max(w["size"] for w in ln)
                bold = all("bold" in w["fontname"].lower() for w in ln)
                if top > page.height * 0.93 and (size < body_size or _FOOTER_PAGE.search(text)):
                    continue  # running footer
                is_title = pno == 1 and size >= max_size - 0.1 and size > body_size * 1.3
                is_heading = bold and len(text.split()) <= 14 and (
                    size >= body_size * 1.1 or _NUMBERED_HEADING.match(text)
                )
                if not seen_section and not is_title and not is_heading and _META_LINE.match(text):
                    continue  # document-control card rendered as plain text rather than a table
                if is_title or is_heading:
                    seen_section = seen_section or (is_heading and not is_title)
                    flush()
                    items.append((top, Block("heading", text, level=0 if is_title else 1, page=pno)))
                    prev_bottom = ln[0]["bottom"]
                    continue
                starts_item = text[:1] in "•-–*" or _NUMBERED_HEADING.match(text) is not None
                gap = (top - prev_bottom) if prev_bottom is not None else 0
                if para and (starts_item or gap > size * 0.9):
                    flush()
                if not para:
                    para_top = top
                para.append(text)
                prev_bottom = ln[0]["bottom"]
            flush()

            items.sort(key=lambda it: it[0])
            blocks.extend(b for _, b in items)
        page_count = len(pdf.pages)
    return ParsedDocument(blocks, page_count=page_count, ocr_pages=ocr_pages)


# -------------------------------------------------------------------------- DOCX
def parse_docx(path: Path) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    blocks: list[Block] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            sl = style.lower()
            m = re.match(r"heading\s*(\d+)", sl)
            runs = [r for r in p.runs if r.text.strip()]
            all_bold = bool(runs) and all(r.bold for r in runs)
            if sl == "title":
                blocks.append(Block("heading", text, level=0))
            elif m:
                blocks.append(Block("heading", text, level=int(m.group(1))))
            elif all_bold and len(text.split()) <= 12 and _NUMBERED_HEADING.match(text):
                blocks.append(Block("heading", text, level=1))
            elif "list" in sl:
                blocks.append(Block("text", "- " + text))
            else:
                blocks.append(Block("text", text))
        elif tag == "tbl":
            t = Table(child, d)
            rows = []
            for r in t.rows:
                cells, last = [], None
                for c in r.cells:  # merged cells repeat; collapse adjacent duplicates
                    if c._tc is last:
                        continue
                    last = c._tc
                    cells.append(_clean_cell(c.text))
                rows.append(cells)
            if rows and not is_metadata_table(rows):
                blocks.append(Block("table", rows=rows))
    return ParsedDocument(blocks)


# -------------------------------------------------------------------------- HTML
def parse_html(path: Path) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    body = soup.body or soup
    wanted = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "dt", "dd", "blockquote", "pre"}
    containers = {"table", "li", "p", "dd", "blockquote", "pre"}
    blocks: list[Block] = []
    for el in body.find_all(list(wanted)):
        if any(a.name in containers for a in el.parents):
            continue
        if el.name == "table":
            rows = [[_clean_cell(c.get_text(" ")) for c in tr.find_all(["td", "th"])] for tr in el.find_all("tr")]
            rows = [r for r in rows if r]
            if rows and not is_metadata_table(rows):
                blocks.append(Block("table", rows=rows))
        elif el.name[0] == "h":
            text = _clean(el.get_text(" "))
            if text:
                lvl = int(el.name[1])
                blocks.append(Block("heading", text, level=0 if lvl == 1 else lvl - 1))
        else:
            text = _clean(el.get_text(" "))
            if text:
                blocks.append(Block("text", ("- " + text) if el.name == "li" else text))
    if not blocks:
        text = _clean(body.get_text(" "))
        if text:
            blocks.append(Block("text", text))
    return ParsedDocument(blocks)


# ---------------------------------------------------------------------- Markdown
def parse_markdown(path: Path) -> ParsedDocument:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    blocks: list[Block] = []
    para: list[str] = []
    table: list[list[str]] = []

    def flush_para():
        nonlocal para
        if para:
            blocks.append(Block("text", _clean(" ".join(para))))
        para = []

    def flush_table():
        nonlocal table
        if table and not is_metadata_table(table):
            blocks.append(Block("table", rows=table))
        table = []

    for raw in lines:
        line = raw.rstrip()
        if line.strip().startswith("|"):
            flush_para()
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # separator row
            table.append(cells)
            continue
        flush_table()
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_para()
            lvl = len(m.group(1))
            blocks.append(Block("heading", _clean(m.group(2)), level=0 if lvl == 1 else lvl - 1))
        elif not line.strip():
            flush_para()
        elif re.match(r"^\s*([-*+]|\d+\.)\s+", line):
            flush_para()
            blocks.append(Block("text", "- " + _clean(re.sub(r"^\s*([-*+]|\d+\.)\s+", "", line))))
        else:
            para.append(line.strip())
    flush_para()
    flush_table()
    return ParsedDocument(blocks)


def parse_file(path: Path, ocr_lang: str = "eng") -> ParsedDocument:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path, ocr_lang)
    if ext == ".docx":
        return parse_docx(path)
    if ext in (".html", ".htm"):
        return parse_html(path)
    if ext in (".md", ".markdown", ".txt"):
        return parse_markdown(path)
    raise ValueError(f"unsupported file type: {ext}")
